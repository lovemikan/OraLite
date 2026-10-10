"""结果集导出：Excel / CSV / JSON / SQL 插入脚本。"""
from __future__ import annotations

import csv
import json
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable, Iterable, List, Optional, Sequence, Tuple

from .validator import format_value, to_excel_value

ILLEGAL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
MAX_SHEET_ROWS = 1_048_000


def _clean(v: Any) -> Any:
    if isinstance(v, str):
        return ILLEGAL_CHARS.sub("", v)
    return v


def _safe_sheet_name(name: str, used: set) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", (name or "数据").strip())[:31] or "数据"
    n, i = base, 1
    while n in used:
        i += 1
        suffix = f"_{i}"
        n = base[: 31 - len(suffix)] + suffix
    used.add(n)
    return n


def export_excel(path: str, columns: Sequence[str], rows: Iterable[Sequence[Any]],
                 sheet_name: str = "数据", progress: Optional[Callable[[int], bool]] = None,
                 title: str = "", date_format: str = "%Y-%m-%d %H:%M:%S") -> Tuple[int, List[str]]:
    """导出为 xlsx，返回 (总行数, sheet 名列表)。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    sheets: List[str] = []
    ws = None
    row_idx = 0
    total = 0
    widths: List[int] = []
    col_count = len(columns)

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="4A6FA5")
    header_align = Alignment(horizontal="center", vertical="center")

    def new_sheet():
        nonlocal ws, row_idx, widths
        nm = _safe_sheet_name(sheet_name, used)
        sheets.append(nm)
        ws = wb.create_sheet(nm)
        ws.append([_clean(c) for c in columns])
        for c in range(1, col_count + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        ws.freeze_panes = "A2"
        if col_count:
            ws.auto_filter.ref = f"A1:{get_column_letter(col_count)}1"
        widths = [min(max(len(str(c)) + 2, 10), 40) for c in columns]
        row_idx = 1

    new_sheet()

    for r_i, row in enumerate(rows):
        if row_idx >= MAX_SHEET_ROWS:
            for c, w in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(c)].width = w
            new_sheet()
        values = []
        for i, v in enumerate(row):
            v = to_excel_value(v)
            if isinstance(v, str):
                v = _clean(v)
            values.append(v)
            if len(values) <= len(widths):
                try:
                    ln = len(format_value(row[i], date_format) if not isinstance(v, str) else v)
                    if ln + 2 > widths[len(values) - 1] and widths[len(values) - 1] < 45:
                        widths[len(values) - 1] = min(ln + 2, 45)
                except Exception:
                    pass
        ws.append(values)
        row_idx += 1
        total += 1
        if progress and total % 1000 == 0:
            if progress(total) is False:
                break

    for c, w in enumerate(widths, start=1):
        try:
            ws.column_dimensions[get_column_letter(c)].width = w
        except Exception:
            pass

    # 日期列格式
    try:
        from openpyxl.styles import Alignment as A

        for ws_ in wb.worksheets:
            for row in ws_.iter_rows(min_row=2, max_row=min(ws_.max_row, 200)):
                for cell in row:
                    if isinstance(cell.value, (datetime, date)):
                        cell.number_format = "yyyy-mm-dd hh:mm:ss"
    except Exception:
        pass

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return total, sheets


def export_csv(path: str, columns: Sequence[str], rows: Iterable[Sequence[Any]],
               progress: Optional[Callable[[int], bool]] = None,
               date_format: str = "%Y-%m-%d %H:%M:%S", encoding: str = "utf-8-sig") -> int:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    total = 0
    with open(path, "w", encoding=encoding, newline="") as f:
        w = csv.writer(f)
        w.writerow(list(columns))
        for row in rows:
            w.writerow([format_value(v, date_format) for v in row])
            total += 1
            if progress and total % 1000 == 0:
                if progress(total) is False:
                    break
    return total


def export_json(path: str, columns: Sequence[str], rows: Iterable[Sequence[Any]],
                progress: Optional[Callable[[int], bool]] = None,
                date_format: str = "%Y-%m-%d %H:%M:%S") -> int:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    total = 0
    with open(path, "w", encoding="utf-8") as f:
        f.write("[")
        first = True
        for row in rows:
            obj = {}
            for c, v in zip(columns, row):
                if v is None:
                    obj[c] = None
                elif isinstance(v, (datetime, date)):
                    obj[c] = v.strftime(date_format)
                elif isinstance(v, Decimal):
                    obj[c] = float(v)
                else:
                    obj[c] = _clean(v) if isinstance(v, str) else v
            f.write(("" if first else ",") + json.dumps(obj, ensure_ascii=False))
            first = False
            total += 1
            if progress and total % 1000 == 0:
                if progress(total) is False:
                    break
        f.write("]")
    return total


def export_insert_sql(path: str, table: str, columns: Sequence[str], rows: Iterable[Sequence[Any]],
                      progress: Optional[Callable[[int], bool]] = None,
                      date_format: str = "%Y-%m-%d %H:%M:%S", batch_commit: int = 100) -> int:
    """导出成可回放的 INSERT 脚本。"""
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    total = 0
    cols = ", ".join(columns)
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"-- 由 Oracle 轻量客户端生成：{table} 共导出数据脚本\n")
        f.write("-- 执行前请确认目标表结构一致\n")
        for row in rows:
            vals = []
            for v in row:
                if v is None:
                    vals.append("NULL")
                elif isinstance(v, (datetime, date)):
                    vals.append(f"TO_DATE('{v.strftime('%Y-%m-%d %H:%M:%S')}','YYYY-MM-DD HH24:MI:SS')")
                elif isinstance(v, (int, float, Decimal)):
                    vals.append(str(v))
                else:
                    vals.append("'" + str(_clean(v)).replace("'", "''") + "'")
            f.write(f"INSERT INTO {table} ({cols}) VALUES ({', '.join(vals)});\n")
            total += 1
            if total % batch_commit == 0:
                f.write("COMMIT;\n")
            if progress and total % 1000 == 0:
                if progress(total) is False:
                    break
        f.write("COMMIT;\n")
    return total


def suggest_filename(prefix: str, ext: str) -> str:
    from datetime import datetime as _dt

    safe = re.sub(r"[\\/:*?\"<>|]", "_", prefix or "数据")[:40]
    return f"{safe}_{_dt.now().strftime('%Y%m%d_%H%M%S')}.{ext.lstrip('.')}"
